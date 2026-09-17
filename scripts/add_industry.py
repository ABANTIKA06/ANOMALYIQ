import csv
import glob
import os
import openpyxl

def main():
    # Load mapping
    mapping = {}
    with open('ind_nifty500list.csv', mode='r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            mapping[row['Symbol']] = row['Industry']
    
    # Process files
    files = glob.glob('screener_exports/*.xlsx')
    total = len(files)
    print(f"Found {total} files.")
    for i, filepath in enumerate(files):
        filename = os.path.basename(filepath)
        symbol = filename.replace('.xlsx', '')
        
        if symbol in mapping:
            industry = mapping[symbol]
            try:
                wb = openpyxl.load_workbook(filepath)
                if 'Data Sheet' in wb.sheetnames:
                    ws = wb['Data Sheet']
                    ws['A10'] = 'Industry'
                    ws['B10'] = industry
                    wb.save(filepath)
                else:
                    print(f"Skipping {symbol}: 'Data Sheet' not found.")
            except Exception as e:
                print(f"Error processing {symbol}: {e}")
        else:
            print(f"Symbol {symbol} not found in mapping.")
            
        if (i + 1) % 50 == 0:
            print(f"Processed {i + 1}/{total}")
            
    print("Done!")

if __name__ == '__main__':
    main()
